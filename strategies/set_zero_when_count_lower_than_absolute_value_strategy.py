from openpyxl.worksheet.worksheet import Worksheet

from strategies.transform_strategy import TransformStrategy


class SetZeroWhenNumberLowerThanAbsoluteValueStrategy(TransformStrategy):
    """
    Strategy that sets to zero the integer and percentage values in pairs of columns when the integer value
    is less than an absolute value.
    """

    __ABSOLUTE_VALUE = 100

    def transform_excel_sheet(self, sheet: Worksheet, start_column_index: int) -> None:
        # Iterate over the rows starting from the second row
        for row in range(2, sheet.max_row + 1):
            col = start_column_index
            while col < sheet.max_column:
                int_value_cell = sheet.cell(row=row, column=col)
                perc_value_cell = sheet.cell(row=row, column=col + 1)

                # Check if the integer value is less than the absolute value
                if int_value_cell.value is not None and int_value_cell.value < self.__ABSOLUTE_VALUE:
                    int_value_cell.value = 0
                    int_value_cell.font = self.red_font  # Set the text color to red
                    perc_value_cell.value = 0
                    perc_value_cell.font = self.red_font  # Set the text color to red

                col += 2  # Move to the next pair of columns

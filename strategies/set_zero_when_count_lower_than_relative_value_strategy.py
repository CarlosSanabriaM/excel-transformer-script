from openpyxl.worksheet.worksheet import Worksheet

from strategies.transform_strategy import TransformStrategy


class SetZeroWhenNumberLowerThanRelativeValueStrategy(TransformStrategy):
    """
    Strategy that sets to zero the integer and percentage values in pairs of columns when the integer value
    relative to the count column (in percentage) is less than a relative value (in percentage).
    """

    __MIN_RELATIVE_VALUE_PERCENTAGE = 0.1

    def transform_excel_sheet(self, sheet: Worksheet, start_column_index: int, count_column_name: str) -> None:
        # Iterate over the rows starting from the second row
        for row in range(2, sheet.max_row + 1):
            # Get the count value for the current row
            count_column_index = self.__get_column_index_from_name(sheet, count_column_name)
            count_value = sheet.cell(row=row, column=count_column_index).value

            col = start_column_index
            while col < sheet.max_column:
                int_value_cell = sheet.cell(row=row, column=col)
                perc_value_cell = sheet.cell(row=row, column=col + 1)

                # Check if the integer value relative to the count column is less than the relative value
                if int_value_cell.value is not None:
                    relative_value_percentage = int_value_cell.value / count_value * 100
                    if relative_value_percentage < self.__MIN_RELATIVE_VALUE_PERCENTAGE:
                        int_value_cell.value = 0
                        int_value_cell.font = self.red_font  # Set the text color to red
                        perc_value_cell.value = 0
                        perc_value_cell.font = self.red_font  # Set the text color to red

                col += 2  # Move to the next pair of columns

    def __get_column_index_from_name(self, sheet: Worksheet, column_name: str):
        for col in range(1, sheet.max_column + 1):
            cell_value = sheet.cell(row=1, column=col).value
            if cell_value == column_name:
                return col
        return None
